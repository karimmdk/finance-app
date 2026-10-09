"""
پارس فایل Excel نمونه بانکی — معادل src/importers/generic-bank-xlsx.ts.
از openpyxl با دسترسی مستقیم به آدرس سلول استفاده می‌شود (نه ردیف‌های آرایه‌ای) — دقیقاً همان
رویکردی که در نسخه Node بعد از پیدا شدن باگ column-shift انتخاب شد؛ openpyxl از ابتدا با آدرس
مطلق سلول کار می‌کند، پس این کلاس مشکل SheetJS (که این پروژه رویش تست و رفع شد) را ندارد.
"""
import io
from openpyxl import load_workbook
from app.lib.jalali import parse_bank_datetime
from app.lib.duplicate_detection import NormalizedTransaction

EXPENSE_TYPES = {"خرید اینترنتی", "خرید از فروشگاه", "قبض", "کارمزد", "انتقال به کارت", "انتقال به سپرده"}
INCOME_TYPES = {"دریافت از کارت", "دریافت از سپرده", "واریز سود بانکی"}
ADJUSTMENT_TYPES = {"اصلاح سند"}


def _is_number(v) -> bool:
    if isinstance(v, bool) or v is None:
        return False
    if isinstance(v, (int, float)):
        return True
    try:
        float(str(v).replace(",", "").strip())
        return True
    except ValueError:
        return False


def parse_generic_bank_xlsx(file_bytes: bytes, account_id: int) -> dict:
    wb = load_workbook(io.BytesIO(file_bytes), data_only=True)
    sheet_name = "transaction" if "transaction" in wb.sheetnames else wb.sheetnames[0]
    ws = wb[sheet_name]

    def val(col: str, row: int):
        return ws[f"{col}{row}"].value

    header_meta = {}
    if val("B", 2): header_meta["previous_balance"] = str(val("B", 2))
    if val("B", 4): header_meta["total_deposit"] = str(val("B", 4))
    if val("B", 6): header_meta["total_withdrawal"] = str(val("B", 6))
    if val("B", 7): header_meta["final_balance"] = str(val("B", 7))
    if val("O", 3): header_meta["account_owner"] = str(val("O", 3))
    if val("O", 7): header_meta["iban"] = str(val("O", 7))

    # پیدا کردن ردیف هدر واقعی به‌صورت پویا: اولین ردیفی که ستون G برابر «نوع تراکنش» است
    header_row = None
    for r in range(1, ws.max_row + 1):
        if val("G", r) == "نوع تراکنش":
            header_row = r
            break
    if header_row is None:
        raise ValueError('ردیف هدر جدول یافت نشد (ستون "نوع تراکنش" پیدا نشد) — فرمت فایل با نمونه بررسی‌شده مطابقت ندارد')

    transactions: list[NormalizedTransaction] = []
    issues: list[dict] = []

    for r in range(header_row + 1, ws.max_row + 1):
        balance = val("B", r)
        withdrawal = val("C", r)
        deposit = val("D", r)
        raw_type = val("G", r)
        description = val("K", r)
        document_number = val("P", r)
        date_raw = val("R", r)
        source_row_number = val("S", r)

        if balance is None and date_raw is None:
            continue
        # ردیف‌های فوتر/سرصفحه صفحه (مثل «تاریخ صدور ...» و «صفحه 1») نه تاریخ دارند نه مبلغ: تراکنش نیستند
        # و نباید هشدار «تاریخ خالی است» بدهند.
        if not date_raw and not _is_number(withdrawal) and not _is_number(deposit):
            continue
        if not date_raw:
            issues.append({"row_number": r, "rowNumber": r, "message": "تاریخ تراکنش خالی است — ردیف نادیده گرفته شد"})
            continue

        withdrawal_num = float(withdrawal) if withdrawal else 0
        deposit_num = float(deposit) if deposit else 0

        if withdrawal_num > 0 and deposit_num > 0:
            issues.append({"row_number": r, "rowNumber": r, "message": "هم برداشت و هم واریز غیرصفر است — ردیف مشکوک، نیاز به بررسی دستی"})
        if withdrawal_num == 0 and deposit_num == 0:
            issues.append({"row_number": r, "rowNumber": r, "message": "مبلغ تراکنش صفر است — ردیف نادیده گرفته شد"})
            continue

        if raw_type in ADJUSTMENT_TYPES:
            txn_type = "adjustment"
            signed_amount = deposit_num if deposit_num > 0 else -withdrawal_num
        elif raw_type in INCOME_TYPES:
            txn_type = "income"
            signed_amount = deposit_num
        elif raw_type in EXPENSE_TYPES:
            txn_type = "expense"
            signed_amount = withdrawal_num
        else:
            txn_type = "income" if deposit_num > 0 else "expense"
            signed_amount = deposit_num if deposit_num > 0 else withdrawal_num
            issues.append({"row_number": r, "rowNumber": r, "message": f'نوع تراکنش ناشناخته: "{raw_type}" — به‌صورت {txn_type} فرض شد، لطفاً بررسی کنید'})

        date_iso, time_str = parse_bank_datetime(str(date_raw))

        amount = int(signed_amount) if txn_type == "adjustment" else int(abs(signed_amount))

        transactions.append(NormalizedTransaction(
            account_id=account_id,
            transaction_date=date_iso,
            transaction_time=time_str,
            amount=amount,
            type=txn_type,
            description=str(description) if description else None,
            document_number=str(document_number) if document_number else None,
            balance_after=int(balance) if balance is not None else None,
            source_row_number=int(source_row_number) if source_row_number is not None else None,
            raw_transaction_type=str(raw_type) if raw_type else None,
        ))

    return {"transactions": transactions, "issues": issues, "header_meta": header_meta}
